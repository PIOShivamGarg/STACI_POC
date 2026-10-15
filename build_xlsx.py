"""Staci - build the Excel workbook (README + results + input data).

Run after forecast.py and risk.py, with the same FORECAST_MODEL.
Writes Staci_8.2_POC_<model>.xlsx.
"""
import json
import os

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

from config import DATA_DIR, MODEL, OUT_DIR

WORKBOOK = f"Staci_8.2_POC_{MODEL}.xlsx"

# sheet name -> source CSV, in the order the sheets appear (results first, then inputs)
RESULT_SHEETS = {
    "inventory_risk": f"{OUT_DIR}/inventory_risk.csv",
    "order_volume_forecast": f"{OUT_DIR}/order_volume_forecast.csv",
    "sku_forecast": f"{OUT_DIR}/sku_forecast.csv",
    "backtest_accuracy": f"{OUT_DIR}/backtest_accuracy.csv",
    "feature_importance": f"{OUT_DIR}/feature_importance.csv",
}
INPUT_SHEETS = {name: f"{DATA_DIR}/{name}.csv"
                for name in ["items", "inventory_snapshot", "inbound_asn", "events", "daily_demand"]}

FONT = Font(name="Arial", size=10)
BOLD_FONT = Font(name="Arial", size=10, bold=True)
HEADER_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
RISK_FILLS = {state: PatternFill("solid", fgColor=color)
              for state, color in {"RED": "F8CBAD", "AMBER": "FFE699", "GREEN": "C6EFCE", "BLUE": "BDD7EE"}.items()}
STYLE_ROW_LIMIT = 5000        # skip per-cell fonts on bigger sheets (daily_demand) - too slow
WIDTH_SAMPLE_ROWS = 200       # rows sampled to size each column
MODEL_NAME = {"lgbm": "LightGBM", "xgboost": "XGBoost", "chronos": "Chronos-2",
              "stats": "Statistical (MSTL+ETS / TSB)"}.get(MODEL, MODEL)
TREE_DESCRIPTION = ("global model trained on lag / rolling / event features: P50 (tweedie objective) "
                    "and P90 (quantile 0.9), recursive 13-week forecast")
MODEL_DESCRIPTION = {
    "stats": ("one small statistical model per SKU (statsforecast, Apache 2.0): regular sellers use MSTL + ETS "
              "(yearly pattern + smoothed level/trend), slow movers (>30% zero weeks) use TSB; "
              "P90 = upper end of the 80% prediction interval; events not used"),
    "chronos": ("pre-trained time-series foundation model (Amazon, Apache 2.0), no training on Staci data: "
                "forecasts the 0.5 and 0.9 quantiles directly from each SKU's weekly history, "
                "with event uplift as a known future input"),
}.get(MODEL, TREE_DESCRIPTION)


def readme_sheet():
    with open(f"{OUT_DIR}/summary.json") as f:
        summary = json.load(f)
    with open(f"{OUT_DIR}/risk_summary.json") as f:
        risk_counts = json.load(f)["risk_counts"]

    def backtest_line(section):
        parts = []
        for window, label in [("last_13_weeks", "last 13 weeks"), ("peak_season_last_year", "last peak season")]:
            b = summary["backtest"][window][section]
            parts.append(f'{label}: WAPE {b["model"]:.1%} (naive {b["naive"]:.1%}), P90 coverage {b["p90_coverage"]:.1%}')
        return "; ".join(parts)

    rows = [
        ("About", "Synthetic, fictional dataset for the Staci Section 8.2 POC. Replace each input sheet "
                  "with a real Staci extract using the same columns."),
        ("As-of date", summary["as_of"]),
        ("Method version", summary["method_version"]),
        ("Horizon", f'{summary["horizon_weeks"]} weeks'),
        ("Model", f"{MODEL_NAME}: {MODEL_DESCRIPTION}"),
        ("Backtest: Order Volume", "Orders per client / warehouse / channel / week - " + backtest_line("order_volume")),
        ("Backtest: Shortfall Risk (proxy)", "Demand during each SKU's lead time - " + backtest_line("shortfall_risk")
                                             + ". A true alert backtest needs historical stock snapshots."),
        ("Backtest: SKU demand", "Units per SKU per week - " + backtest_line("sku_demand")),
        ("Risk counts", ", ".join(f"{state} {n}" for state, n in risk_counts.items())),
        ("INPUT items", "Item master: SKU, client, warehouse, channel, supplier lead time (+ variability), "
                        "safety-stock days, units per order"),
        ("INPUT daily_demand", "Daily shipped units and order counts per SKU / warehouse / channel (3 years)"),
        ("INPUT inventory_snapshot", "On-hand, allocated and on-hold qty at the as-of date"),
        ("INPUT inbound_asn", "Open / in-transit inbound shipments with expected arrival date and qty"),
        ("INPUT events", "Known promotions and peak periods (Black Friday etc.)"),
        ("OUTPUT inventory_risk", "One row per SKU: days of cover, projected stockout date (P50 / P90), "
                                  "RED/AMBER/GREEN/BLUE state, reason, suggested reorder qty"),
        ("OUTPUT order_volume_forecast", "Weekly units and orders (P50 / P90) per client / warehouse / channel "
                                         "- for labour and capacity planning"),
        ("OUTPUT sku_forecast", "Weekly SKU demand P50 / P90 for next 13 weeks, with created_at and "
                                "method_version for audit"),
        ("OUTPUT backtest_accuracy", f"Per-SKU forecast vs actual for two historical windows, {MODEL_NAME} vs naive"),
        ("OUTPUT feature_importance", f"Which inputs drive the {MODEL_NAME} P50 forecast (share of total gain); "
                                      "tree models only"),
        ("Risk states", "RED = runs out before stock can arrive; AMBER = reorder now; GREEN = covered; "
                        "BLUE = excess (>90 days cover)"),
    ]
    return pd.DataFrame(rows, columns=["Item", "Detail"])


def style_sheet(ws):
    if ws.max_row <= STYLE_ROW_LIMIT:
        for row in ws.iter_rows():
            for cell in row:
                cell.font = FONT
    for cell in ws[1]:
        cell.font, cell.fill = HEADER_FONT, HEADER_FILL
    ws.freeze_panes = "A2"
    max_width = 90 if ws.title == "README" else 40
    for col in ws.columns:
        longest = max(len(str(cell.value or "")) for cell in list(col)[:WIDTH_SAMPLE_ROWS])
        ws.column_dimensions[col[0].column_letter].width = min(max(10, longest + 2), max_width)


def color_risk_states(ws):
    col = [cell.value for cell in ws[1]].index("risk_state") + 1
    for r in range(2, ws.max_row + 1):
        cell = ws.cell(r, col)
        cell.fill, cell.font = RISK_FILLS[cell.value], BOLD_FONT


def main():
    sheets = {"README": readme_sheet()}
    sheets.update({name: pd.read_csv(path) for name, path in {**RESULT_SHEETS, **INPUT_SHEETS}.items()
                   if os.path.exists(path)})   # feature_importance is missing for chronos
    with pd.ExcelWriter(WORKBOOK, engine="openpyxl") as writer:
        for name, df in sheets.items():
            df.to_excel(writer, sheet_name=name, index=False)

    wb = load_workbook(WORKBOOK)
    for ws in wb:
        style_sheet(ws)
    color_risk_states(wb["inventory_risk"])
    wb.save(WORKBOOK)
    print(f"Wrote {WORKBOOK}")


if __name__ == "__main__":
    main()
